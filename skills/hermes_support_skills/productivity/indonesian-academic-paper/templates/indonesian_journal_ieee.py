"""
Template: Indonesian Academic Journal Paper (IEEE Style)
========================================================
Key format rules:
- Title: UPPERCASE, bold, centered, 14pt
- "Abstrak:" bold (Indonesian) + "Abstract:" bold italic (English)
- Numbered sections: 1. PENDAHULUAN, 2. METODE, 3. HASIL DAN PEMBAHASAN, 4. KESIMPULAN
- References: IEEE numbered [1], [2], [3]
- Tables: BLACK AND WHITE only, single border, no color
- Font: Times New Roman 12pt, Line spacing: Pt(20), Margins: 3cm/2.5cm
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def create_indonesian_journal():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(3); s.bottom_margin = Cm(3)
        s.left_margin = Cm(3); s.right_margin = Cm(2.5)
    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'; st.font.size = Pt(12)
    st.paragraph_format.line_spacing = Pt(20)
    st.paragraph_format.space_after = Pt(0); st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return doc

def h(doc, text, level=1):
    hd = doc.add_heading(text, level=level)
    for r in hd.runs:
        r.font.name = 'Times New Roman'; r.font.color.rgb = RGBColor(0,0,0); r.font.size = Pt(12)
    hd.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hd.paragraph_format.space_before = Pt(12); hd.paragraph_format.space_after = Pt(6)

def p(doc, text, bold=False):
    pa = doc.add_paragraph(); r = pa.add_run(text)
    r.font.name = 'Times New Roman'; r.font.size = Pt(12); r.bold = bold
    pa.paragraph_format.line_spacing = Pt(20); pa.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

def tbl_bw(doc, headers, rows):
    t = doc.add_table(rows=1+len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = True
    tbl = t._tbl; tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement('w:tblPr')
    borders = OxmlElement('w:tblBorders')
    for bn in ['top','left','bottom','right','insideH','insideV']:
        b = OxmlElement(f'w:{bn}'); b.set(qn('w:val'),'single')
        b.set(qn('w:sz'),'4'); b.set(qn('w:space'),'0'); b.set(qn('w:color'),'000000')
        borders.append(b)
    tblPr.append(borders)
    for i, hd in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = hd
        for pa in c.paragraphs: pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in c.paragraphs[0].runs: r.bold=True; r.font.name='Times New Roman'; r.font.size=Pt(10)
    for ri, row in enumerate(rows):
        for ci, v in enumerate(row):
            c = t.rows[ri+1].cells[ci]; c.text = str(v)
            for pa in c.paragraphs: pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in pa.runs: r.font.name='Times New Roman'; r.font.size=Pt(10)
    doc.add_paragraph()

# Example usage:
if __name__ == '__main__':
    doc = create_indonesian_journal()
    pa = doc.add_paragraph(); pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = pa.add_run('JUDUL ARTIKEL DALAM HURUF KAPITAL')
    r.font.name='Times New Roman'; r.font.size=Pt(14); r.bold=True
    doc.add_paragraph()
    
    # Indonesian abstract
    pa = doc.add_paragraph()
    r = pa.add_run('Abstrak: '); r.bold=True; r.font.name='Times New Roman'; r.font.size=Pt(12)
    r = pa.add_run('Abstrak dalam Bahasa Indonesia...'); r.font.name='Times New Roman'; r.font.size=Pt(12)
    pa = doc.add_paragraph()
    r = pa.add_run('Kata kunci: '); r.bold=True; r.font.name='Times New Roman'; r.font.size=Pt(12)
    r = pa.add_run('Kata1, Kata2, Kata3'); r.font.name='Times New Roman'; r.font.size=Pt(12); r.italic=True
    
    # English abstract
    pa = doc.add_paragraph()
    r = pa.add_run('Abstract: '); r.bold=True; r.font.name='Times New Roman'; r.font.size=Pt(12); r.italic=True
    r = pa.add_run('Abstract in English...'); r.font.name='Times New Roman'; r.font.size=Pt(12); r.italic=True
    pa = doc.add_paragraph()
    r = pa.add_run('Keywords: '); r.bold=True; r.font.name='Times New Roman'; r.font.size=Pt(12); r.italic=True
    r = pa.add_run('Keyword1, Keyword2, Keyword3'); r.font.name='Times New Roman'; r.font.size=Pt(12); r.italic=True
    
    # Sections
    h(doc, '1. PENDAHULUAN')
    p(doc, 'Isi pendahuluan dengan referensi [1], [2]...')
    h(doc, '2. METODE')
    p(doc, 'Deskripsi metode penelitian...')
    h(doc, '3. HASIL DAN PEMBAHASAN')
    p(doc, 'Hasil penelitian...')
    h(doc, '4. KESIMPULAN')
    p(doc, 'Kesimpulan...')
    
    # IEEE References
    h(doc, 'DAFTAR PUSTAKA')
    refs = [
        '[1] J. M. Wing, "Computational thinking," Communications of the ACM, vol. 49, no. 3, pp. 33–35, 2006.',
        '[2] K. Brennan and M. Resnick, "New frameworks for studying CT," Proceedings 2012, pp. 1–25, 2012.',
    ]
    for ref in refs:
        pa = doc.add_paragraph(); r = pa.add_run(ref)
        r.font.name='Times New Roman'; r.font.size=Pt(12)
        pa.paragraph_format.line_spacing=Pt(20); pa.paragraph_format.space_after=Pt(4)
    
    doc.save('template_jurnal_indonesia.docx')
    print("Template saved!")
