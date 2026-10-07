import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os

doc = docx.Document()

# Set standard margins (1 inch)
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Helper for cell background color
def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

# Helper for cell margins
def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

# Title Block
title_p = doc.add_paragraph()
title_p.paragraph_format.space_before = Pt(0)
title_p.paragraph_format.space_after = Pt(4)
title_run = title_p.add_run("PRODUCT REQUIREMENT DOCUMENT (PRD)")
title_run.font.name = "Arial"
title_run.font.size = Pt(22)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900

sub_p = doc.add_paragraph()
sub_p.paragraph_format.space_after = Pt(14)
sub_run = sub_p.add_run("Widya.X OS - Unified LMS, Multi-Branch ERP and Marketing Operations Platform")
sub_run.font.name = "Arial"
sub_run.font.size = Pt(13)
sub_run.font.bold = True
sub_run.font.color.rgb = RGBColor(220, 38, 38) # Red 600

# Meta Card Box
meta_table = doc.add_table(rows=4, cols=2)
meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
meta_data = [
    ("Lembaga / Institusi:", "Widya.X (Robotics, AI, Agentic and Drone Academy)"),
    ("Operational Manager / Penanggung Jawab:", "Andi Saputra"),
    ("Fokus 4 Pilar Kurikulum:", "Robotika Humanoid/Embedded, Artificial Intelligence, AI Agentic, dan Drone FPV/UAV"),
    ("Cakupan Sistem:", "Manajemen 2 Cabang, Siswa, Guru/Mentor, LMS 4 Pilar, Digital Marketing CRM dan Billing QRIS")
]
for i, (k, v) in enumerate(meta_data):
    row = meta_table.rows[i]
    cell_k, cell_v = row.cells[0], row.cells[1]
    cell_k.width = Inches(2.2)
    cell_v.width = Inches(4.3)
    set_cell_background(cell_k, "F1F5F9")
    set_cell_background(cell_v, "FFFFFF")
    set_cell_margins(cell_k, 80, 80, 100, 100)
    set_cell_margins(cell_v, 80, 80, 100, 100)
    
    pk = cell_k.paragraphs[0]
    pk.paragraph_format.space_after = Pt(0)
    rk = pk.add_run(k)
    rk.font.name = "Arial"
    rk.font.size = Pt(9.5)
    rk.font.bold = True
    rk.font.color.rgb = RGBColor(51, 65, 85)
    
    pv = cell_v.paragraphs[0]
    pv.paragraph_format.space_after = Pt(0)
    rv = pv.add_run(v)
    rv.font.name = "Arial"
    rv.font.size = Pt(9.5)
    rv.font.color.rgb = RGBColor(15, 23, 42)

doc.add_paragraph().paragraph_format.space_after = Pt(10)

def add_heading_1(text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    r = h.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)

def add_heading_2(text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    r = h.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(11.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 41, 59)

def add_body(text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        rb = p.add_run(bold_prefix + " ")
        rb.font.name = "Calibri"
        rb.font.size = Pt(10.5)
        rb.font.bold = True
        rb.font.color.rgb = RGBColor(15, 23, 42)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.italic = italic
    r.font.color.rgb = RGBColor(51, 65, 85)

def add_bullet(bold_label, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    rb = p.add_run(bold_label + ": ")
    rb.font.name = "Calibri"
    rb.font.size = Pt(10.5)
    rb.font.bold = True
    rb.font.color.rgb = RGBColor(15, 23, 42)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(51, 65, 85)

# SECTION 1
add_heading_1("1. Ringkasan Eksekutif dan Visi Produk")
add_body("Widya.X OS adalah platform terpadu (Single Source of Truth) berbasis web responsive yang dirancang untuk mengotomasi dan mensentralisasi seluruh proses operasional, manajemen akademik, pemantauan dua cabang, distribusi kurikulum 4 pilar, serta otomatisasi akuisisi prospek pemasaran digital untuk lembaga les teknologi Widya.X.")
add_body("Platform ini memecahkan tantangan fragmentasi data antar-cabang, pencatatan presensi manual, keterlambatan pelaporan progres ke orang tua, dan ketidakefisienan tindak lanjut calon siswa baru.")

# SECTION 2
add_heading_1("2. Struktur Peran Pengguna (Role-Based Access Control)")
add_bullet("Super Admin (Operational Manager)", "Akses penuh lintas 2 cabang, visualisasi P dan L finansial, performa mentor, analitik akuisisi leads, dan kontrol mutasi inventaris hardware.")
add_bullet("Branch Admin (Front Office Cabang 1 dan 2)", "Pendaftaran siswa baru cabang, operasional kasir dan verifikasi SPP, jadwal ruang lab, dan interaksi layanan orang tua.")
add_bullet("Instruktur / Mentor Lab", "Presensi kelas digital, pengisian jurnal mengajar harian (lock-session), penilaian rapor berkala, serta akses modul dan kode proyek.")
add_bullet("Student and Parent Portal (PWA)", "Jadwal sesi belajar, repositori modul dan video tutorial, galeri portofolio karya robotik/drone siswa, rapor digital, dan invoice pembayaran QRIS.")

# SECTION 3
add_heading_1("3. Rincian 6 Modul Fungsional Platform")

add_heading_2("Modul 1: Multi-Branch and Operational Management (2 Cabang)")
add_bullet("Branch Switcher", "Kemampuan beralih data instan antara Cabang 1 (Pusat), Cabang 2 (Satelit), atau tampilan konsolidasi seluruh cabang.")
add_bullet("Room and Arena Scheduler", "Penjadwalan presisi pemakaian Lab Robotika, Lab Komputasi AI, dan Arena Jaring Terbang Drone guna mencegah bentrok jadwal.")
add_bullet("Hardware Asset Tracking", "Pelacakan inventaris unit robot (UBTECH Yanshee, uKit), mikrokontroler (Arduino/ESP32), drone FPV, transmitter, dan status pemeliharaan.")

add_heading_2("Modul 2: Student Management and Academic Lifecycle")
add_bullet("Database Siswa Terpadu", "Profil lengkap, kontak darurat wali murid, asal sekolah, dan rekam jejak level pembelajaran.")
add_bullet("Smart Attendance", "Presensi digital sekali klik oleh mentor atau scan QR kartu siswa yang otomatis memotong kuota sesi bulanan.")
add_bullet("Student Portfolio Hub", "Galeri karya digital yang mendokumentasikan foto, video terbang drone, dan kode pemrograman robot buatan siswa.")
add_bullet("E-Report Card 5 Dimensi", "Rapor berkala 3 bulanan berstandar industri: Logika/Problem Solving (25%), Keterampilan Hardware (25%), Kreativitas (20%), Resiliensi Troubleshooting (15%), dan Komunikasi (15%).")

add_heading_2("Modul 3: Instructor and Teaching Governance")
add_bullet("Teaching Schedule and Lock Session", "Jadwal mengajar terstruktur dan form jurnal mengajar wajib diisi sebelum sesi kelas dapat ditutup.")
add_bullet("KPI and Kinerja Mentor", "Evaluasi berdasarkan kedisiplinan hadir, ulasan kepuasan wali murid, dan kelengkapan dokumentasi proyek.")
add_bullet("Automated Payroll Calculator", "Kalkulasi otomatis honor mengajar berdasarkan sesi reguler, private VIP, maupun workshop yang telah diselesaikan.")

add_heading_2("Modul 4: Content Management and 4-Pillar LMS")
add_bullet("Pilar 1 - Robotics", "Modul PDF, skema wiring Fritzing, script Arduino/MicroPython, dan dokumentasi YanAPI Yanshee.")
add_bullet("Pilar 2 - Artificial Intelligence", "Notebook OpenCV, dataset klasifikasi gambar, dan script Computer Vision YOLOv8.")
add_bullet("Pilar 3 - AI Agentic", "Workflow multi-agent otonom, template system prompt, dan script function calling Hermes/LangChain.")
add_bullet("Pilar 4 - Drone and Avionics", "Diagram wiring flight controller, konfigurasi Betaflight CLI, dan peta jalur ArduPilot Mission Planner.")

add_heading_2("Modul 5: Digital Marketing CRM and WhatsApp Automation")
add_bullet("Kanban Lead Funnel", "Pelacakan visual prospek: Lead Masuk -> Follow-up -> Jadwal Free Trial -> Hadir Trial -> Siswa Resmi -> Lost.")
add_bullet("WhatsApp Gateway Engine", "Notifikasi otomatis pesan selamat datang, konfirmasi jadwal Free Trial 1 Jam, pengingat H-1 kelas, dan broadcast tagihan SPP.")
add_bullet("Attribution Tracker", "Pelacakan sumber konversi siswa baru (Instagram Ads, TikTok, Roadshow Sekolah, atau Referral).")

add_heading_2("Modul 6: Billing, Payment Gateway and Financial Reports")
add_bullet("Invoice Otomatis SPP", "Penerbitan tagihan digital otomatis setiap tanggal 25 per bulan untuk periode belajar berikutnya.")
add_bullet("Dynamic QRIS 24/7", "Integrasi pembayaran instan m-banking dan e-wallet nasional dengan verifikasi webhook otomatis.")
add_bullet("Executive Financial Summary", "Laporan arus kas masuk/keluar, profit/loss per cabang, dan neraca operasional untuk Operational Manager.")

# SECTION 4
add_heading_1("4. Arsitektur Teknis dan Spesifikasi Stack Rekomendasi")

tech_table = doc.add_table(rows=7, cols=3)
tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Komponen", "Teknologi Terpilih", "Alasan dan Keunggulan"]
for j, h in enumerate(headers):
    c = tech_table.rows[0].cells[j]
    set_cell_background(c, "0F172A")
    set_cell_margins(c, 100, 100, 120, 120)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(h)
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)

tech_data = [
    ("Backend API Engine", "FastAPI (Python 3.11+)", "Performa tinggi asinkron, hemat RAM, integrasi mulus ke AI Agent."),
    ("Frontend UI Dashboard", "Tailwind CSS + Alpine.js / Vue 3", "Antarmuka modern, Clean Light/Dark Theme, cepat dan mobile-first."),
    ("Database Relasional", "PostgreSQL / SQLite Production", "Integritas data multi-cabang terjamin, stabil dan ACID-compliant."),
    ("WhatsApp Service", "WhatsApp Gateway API + Webhook", "Saluran komunikasi nomor satu paling efektif untuk orang tua di Indonesia."),
    ("Payment Gateway", "Dynamic QRIS (Open API)", "Pembayaran instan 24 jam dengan verifikasi otomatis tanpa bukti transfer manual."),
    ("Keamanan dan Server", "JWT Auth, HTTPS Lets Encrypt, Nginx", "Isolasi data sensitif siswa dan sistem keuangan lembaga.")
]

for i, (col1, col2, col3) in enumerate(tech_data):
    row = tech_table.rows[i+1]
    bg = "F8FAFC" if i % 2 == 0 else "FFFFFF"
    for j, val in enumerate([col1, col2, col3]):
        c = row.cells[j]
        set_cell_background(c, bg)
        set_cell_margins(c, 80, 80, 100, 100)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(val)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        if j == 0:
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)
        else:
            r.font.color.rgb = RGBColor(51, 65, 85)

# SECTION 5
add_heading_1("5. Roadmap Implementasi Pengembangan (8 Minggu)")
add_bullet("Phase 1 (Minggu 1 - 3): Core Multi-Branch and Academic CRM", "Desain skema database, manajemen 2 cabang, data master siswa, data guru, dan kalender penjadwalan kelas.")
add_bullet("Phase 2 (Minggu 4 - 6): LMS 4 Pilar and WhatsApp Automation", "Repositori materi robotik/AI/drone, presensi digital, rapor 5 dimensi, dan integrasi broadcast WhatsApp.")
add_bullet("Phase 3 (Minggu 7 - 8): Billing QRIS and Executive Marketing Analytics", "Integrasi payment gateway QRIS, dashboard analitik prospek marketing, pengujian integrasi (UAT), dan peluncuran resmi.")

out_path = "/home/ubuntu/ObsidianVault/Widya_X/PRD_Widya_X_Learning_Management_Platform.docx"
doc.save(out_path)
print("SUCCESSFULLY GENERATED DOCX:", out_path, "Size:", os.path.getsize(out_path))
