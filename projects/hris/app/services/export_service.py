"""
HRIS Export Service
Comprehensive PDF and Excel generation for all HRIS modules.
"""

import os
import sqlite3
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.database import DATABASE_PATH
DB_PATH = DATABASE_PATH

HEADER_FILL = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
ALT_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

CURRENCY_FORMAT = "#,##0"

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _get_db(db=None):
    """Return a SQLite connection. If *db* is a path string, open it."""
    if db is None:
        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        return conn, True
    if isinstance(db, (str,)):
        conn = sqlite3.connect(db)
        conn.row_factory = sqlite3.Row
        return conn, True
    return db, False


def _close_db(conn, opened_here):
    if opened_here and conn:
        conn.close()


def _where_clause(filters, extra_wheres=None):
    """Build WHERE clause from filters dict."""
    wheres = list(extra_wheres or [])
    params = []
    if filters:
        for key, val in filters.items():
            if val is None or val == "":
                continue
            if key == "start_date":
                wheres.append("a.date >= ?")
                params.append(val)
            elif key == "end_date":
                wheres.append("a.date <= ?")
                params.append(val)
            elif key == "period":
                # Period format: "2026-09"
                parts = val.split("-")
                if len(parts) == 2:
                    wheres.append("p.period_year = ? AND p.period_month = ?")
                    params.extend([int(parts[0]), int(parts[1])])
            else:
                wheres.append(f"{key} = ?")
                params.append(val)
    where = f"WHERE {' AND '.join(wheres)}" if wheres else ""
    return where, params


# ---------------------------------------------------------------------------
# Excel styling
# ---------------------------------------------------------------------------

def _style_excel_sheet(ws, headers):
    """Apply professional styling: header, auto-width, freeze panes."""
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    ws.freeze_panes = "A2"

    for col_idx in range(1, len(headers) + 1):
        max_length = len(str(headers[col_idx - 1]))
        for row_idx in range(2, ws.max_row + 1):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is not None:
                max_length = max(max_length, len(str(val)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 4, 50)


def _apply_row_stripes(ws):
    for row_idx in range(2, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = THIN_BORDER
            if row_idx % 2 == 0:
                cell.fill = ALT_FILL


def _build_workbook(headers, rows, currency_cols=None):
    """Create a styled workbook from headers and row lists. Returns bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    for col_idx, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_idx, value=header)

    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if currency_cols and col_idx in currency_cols:
                if value is not None:
                    cell.number_format = CURRENCY_FORMAT

    _style_excel_sheet(ws, headers)
    _apply_row_stripes(ws)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF helpers
# ---------------------------------------------------------------------------

def _build_pdf(data_rows, headers, title="Report", landscape_mode=False):
    """Build a PDF report. Returns bytes."""
    buf = BytesIO()
    page_size = landscape(A4) if landscape_mode else A4
    doc = SimpleDocTemplate(
        buf, pagesize=page_size,
        leftMargin=30, rightMargin=30, topMargin=40, bottomMargin=40,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CustomTitle", parent=styles["Title"],
        fontSize=16, spaceAfter=4, textColor=colors.HexColor("#003366"),
    )
    sub_style = ParagraphStyle(
        "CustomSub", parent=styles["Normal"],
        fontSize=10, spaceAfter=2, textColor=colors.HexColor("#666666"),
    )
    date_style = ParagraphStyle(
        "DateStyle", parent=styles["Normal"],
        fontSize=8, spaceAfter=12, textColor=colors.HexColor("#999999"),
    )

    elements = [
        Paragraph("HRIS SYSTEM", title_style),
        Paragraph(title, sub_style),
        Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", date_style),
        Spacer(1, 12),
    ]

    table_data = [headers]
    for row in data_rows:
        table_data.append([str(v) if v is not None else "" for v in row])

    if not data_rows:
        table_data.append(["No data available"] + [""] * (len(headers) - 1))

    col_count = len(headers)
    page_width = page_size[0] - 60
    col_width = page_width / col_count

    table = Table(table_data, colWidths=[col_width] * col_count, repeatRows=1)
    style_cmds = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#003366")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 1), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]
    # Alternate row colors
    for i in range(2, len(table_data) + 1, 2):
        style_cmds.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#F2F2F2")))

    table.setStyle(TableStyle(style_cmds))
    elements.append(table)

    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()


# ===========================================================================
# EMPLOYEE EXPORTS
# ===========================================================================

EMPLOYEE_SELECT = """
    SELECT
        e.id,
        e.employee_id_str AS employee_id,
        e.full_name,
        e.email,
        e.phone,
        COALESCE(d.name, '') AS department,
        COALESCE(p.title, '') AS position,
        e.hire_date,
        e.status,
        p.salary_max AS salary
    FROM employees e
    LEFT JOIN departments d ON e.department_id = d.id
    LEFT JOIN positions p ON e.position_id = p.id
"""

EMPLOYEE_HEADERS = [
    "ID", "Employee ID", "Full Name", "Email", "Phone",
    "Department", "Position", "Hire Date", "Status", "Salary",
]


def export_employees_excel(db=None, filters=None):
    """Export employees to Excel. Returns bytes."""
    conn, opened = _get_db(db)
    where, params = _where_clause(filters)
    query = f"{EMPLOYEE_SELECT} {where} ORDER BY e.id"
    raw = conn.execute(query, params).fetchall()
    _close_db(conn, opened)
    rows = [
        [r["id"], r["employee_id"], r["full_name"], r["email"], r["phone"],
         r["department"], r["position"], r["hire_date"], r["status"], r["salary"]]
        for r in raw
    ]
    return _build_workbook(EMPLOYEE_HEADERS, rows, currency_cols={10})


def export_employees_pdf(db=None, filters=None):
    """Export employees to PDF. Returns bytes."""
    conn, opened = _get_db(db)
    where, params = _where_clause(filters)
    query = f"{EMPLOYEE_SELECT} {where} ORDER BY e.id"
    raw = conn.execute(query, params).fetchall()
    _close_db(conn, opened)
    rows = [
        [r["id"], r["employee_id"], r["full_name"], r["email"], r["phone"],
         r["department"], r["position"], r["hire_date"], r["status"],
         f'{r["salary"]:,.0f}' if r["salary"] else ""]
        for r in raw
    ]
    return _build_pdf(rows, EMPLOYEE_HEADERS, title="Employee Report", landscape_mode=True)


# ===========================================================================
# ATTENDANCE EXPORTS
# ===========================================================================

ATTENDANCE_SELECT = """
    SELECT
        a.date,
        COALESCE(e.full_name, '') AS employee,
        a.clock_in,
        a.clock_out,
        a.status,
        CASE
            WHEN a.clock_in IS NOT NULL AND a.clock_out IS NOT NULL
            THEN ROUND((JULIANDAY(a.clock_out) - JULIANDAY(a.clock_in)) * 24, 2)
            ELSE NULL
        END AS duration
    FROM attendance a
    LEFT JOIN employees e ON a.employee_id = e.id
"""

ATTENDANCE_HEADERS = ["Date", "Employee", "Clock In", "Clock Out", "Status", "Duration"]


def export_attendance_excel(db=None, filters=None):
    """Export attendance to Excel. Returns bytes."""
    conn, opened = _get_db(db)
    where, params = _where_clause(filters)
    query = f"{ATTENDANCE_SELECT} {where} ORDER BY a.date DESC"
    raw = conn.execute(query, params).fetchall()
    _close_db(conn, opened)
    rows = [
        [r["date"], r["employee"], r["clock_in"], r["clock_out"],
         r["status"], r["duration"]]
        for r in raw
    ]
    return _build_workbook(ATTENDANCE_HEADERS, rows)


def export_attendance_pdf(db=None, filters=None):
    """Export attendance to PDF. Returns bytes."""
    conn, opened = _get_db(db)
    where, params = _where_clause(filters)
    query = f"{ATTENDANCE_SELECT} {where} ORDER BY a.date DESC"
    raw = conn.execute(query, params).fetchall()
    _close_db(conn, opened)
    rows = [
        [r["date"], r["employee"], r["clock_in"], r["clock_out"],
         r["status"], r["duration"]]
        for r in raw
    ]
    return _build_pdf(rows, ATTENDANCE_HEADERS, title="Attendance Report", landscape_mode=True)


# ===========================================================================
# PAYROLL EXPORTS
# ===========================================================================

PAYROLL_SELECT = """
    SELECT
        p.period_month || '/' || p.period_year AS period,
        COALESCE(e.full_name, '') AS employee,
        p.base_salary,
        p.allowance,
        p.overtime_pay AS overtime,
        p.bonus,
        p.deduction AS deductions,
        p.tax_pph21 AS tax,
        p.bpjs_ketenagakerjaan AS bpjs_tk,
        p.bpjs_kesehatan AS bpjs_kes,
        p.net_salary,
        p.status
    FROM payroll p
    LEFT JOIN employees e ON p.employee_id = e.id
"""

PAYROLL_HEADERS = [
    "Period", "Employee", "Base Salary", "Allowance", "Overtime",
    "Bonus", "Deductions", "Tax", "BPJS TK", "BPJS Kes",
    "Net Salary", "Status",
]

PAYROLL_CURRENCY_COLS = {3, 4, 5, 6, 7, 8, 9, 10, 11}


def export_payroll_excel(db=None, filters=None):
    """Export payroll to Excel. Returns bytes."""
    conn, opened = _get_db(db)
    where, params = _where_clause(filters)
    query = f"{PAYROLL_SELECT} {where} ORDER BY p.period_year DESC, p.period_month DESC"
    raw = conn.execute(query, params).fetchall()
    _close_db(conn, opened)
    rows = [
        [r["period"], r["employee"], r["base_salary"], r["allowance"],
         r["overtime"], r["bonus"], r["deductions"], r["tax"],
         r["bpjs_tk"], r["bpjs_kes"], r["net_salary"], r["status"]]
        for r in raw
    ]
    return _build_workbook(PAYROLL_HEADERS, rows, currency_cols=PAYROLL_CURRENCY_COLS)


def export_payroll_pdf(db=None, filters=None):
    """Export payroll to PDF. Returns bytes."""
    conn, opened = _get_db(db)
    where, params = _where_clause(filters)
    query = f"{PAYROLL_SELECT} {where} ORDER BY p.period_year DESC, p.period_month DESC"
    raw = conn.execute(query, params).fetchall()
    _close_db(conn, opened)
    rows = []
    for r in raw:
        rows.append([
            r["period"], r["employee"],
            f'{r["base_salary"]:,.0f}' if r["base_salary"] else "0",
            f'{r["allowance"]:,.0f}' if r["allowance"] else "0",
            f'{r["overtime"]:,.0f}' if r["overtime"] else "0",
            f'{r["bonus"]:,.0f}' if r["bonus"] else "0",
            f'{r["deductions"]:,.0f}' if r["deductions"] else "0",
            f'{r["tax"]:,.0f}' if r["tax"] else "0",
            f'{r["bpjs_tk"]:,.0f}' if r["bpjs_tk"] else "0",
            f'{r["bpjs_kes"]:,.0f}' if r["bpjs_kes"] else "0",
            f'{r["net_salary"]:,.0f}' if r["net_salary"] else "0",
            r["status"],
        ])
    return _build_pdf(rows, PAYROLL_HEADERS, title="Payroll Report", landscape_mode=True)


def export_single_payslip_pdf(db=None, emp_id=1, period_month=None, period_year=None):
    """Generate a single employee payslip PDF. Returns bytes."""
    conn, opened = _get_db(db)
    now = datetime.now()
    m = int(period_month) if period_month else now.month
    y = int(period_year) if period_year else now.year

    query = """
        SELECT p.*, e.full_name, e.employee_id_str, d.name as department_name, pos.title as position_title
        FROM employees e
        LEFT JOIN payroll p ON p.employee_id = e.id AND p.period_month = ? AND p.period_year = ?
        LEFT JOIN departments d ON e.department_id = d.id
        LEFT JOIN positions pos ON e.position_id = pos.id
        WHERE e.id = ?
    """
    row = conn.execute(query, (m, y, emp_id)).fetchone()
    _close_db(conn, opened)

    full_name = row["full_name"] if row and row["full_name"] else f"Employee #{emp_id}"
    emp_code = row["employee_id_str"] if row and row["employee_id_str"] else f"MVP{emp_id:04d}"
    dept = row["department_name"] if row and row["department_name"] else "General"
    pos = row["position_title"] if row and row["position_title"] else "Staff"
    base = row["base_salary"] if row and row["base_salary"] else 7500000
    allowance = row["allowance"] if row and row["allowance"] else int(base * 0.15)
    deduction = row["deduction"] if row and row["deduction"] else int(base * 0.05)
    net = row["net_salary"] if row and row["net_salary"] else (base + allowance - deduction)

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle("TitleStyle", parent=styles["Heading1"], alignment=1, fontSize=16, leading=20, textColor=colors.HexColor("#1e3a5f"))
    subtitle_style = ParagraphStyle("SubTitleStyle", parent=styles["Normal"], alignment=1, fontSize=10, textColor=colors.HexColor("#6c757d"))
    normal_style = styles["Normal"]
    bold_style = ParagraphStyle("BoldStyle", parent=styles["Normal"], fontName="Helvetica-Bold")
    right_normal_style = ParagraphStyle("RightNormal", parent=styles["Normal"], alignment=2)
    right_bold_style = ParagraphStyle("RightBold", parent=styles["Normal"], fontName="Helvetica-Bold", alignment=2)
    header_right_style = ParagraphStyle("HeaderRight", parent=styles["Normal"], fontName="Helvetica-Bold", textColor=colors.white, alignment=2)
    header_left_style = ParagraphStyle("HeaderLeft", parent=styles["Normal"], fontName="Helvetica-Bold", textColor=colors.white, alignment=0)
    net_pay_style = ParagraphStyle("NetPayRight", parent=styles["Normal"], fontName="Helvetica-Bold", alignment=2)
    
    elements = []
    logo_path = "/home/ubuntu/hris/static/img/logo.png"
    if os.path.exists(logo_path):
        from reportlab.platypus import Image as RLImage
        elements.append(RLImage(logo_path, width=45, height=45))
        elements.append(Spacer(1, 4))
    elements.append(Paragraph("<b>PT MITSINDO</b>", title_style))
    elements.append(Paragraph("Human Resource Information System &bull; Confidential Payslip", subtitle_style))
    elements.append(Spacer(1, 15))
    
    info_data = [
        [Paragraph(f"<b>Employee Name:</b> {full_name}", normal_style), Paragraph(f"<b>Period:</b> {m:02d}/{y}", normal_style)],
        [Paragraph(f"<b>Employee ID:</b> {emp_code}", normal_style), Paragraph(f"<b>Department:</b> {dept}", normal_style)],
        [Paragraph(f"<b>Position:</b> {pos}", normal_style), Paragraph(f"<b>Date Generated:</b> {now.strftime('%d-%m-%Y')}", normal_style)],
    ]
    info_table = Table(info_data, colWidths=[260, 260])
    info_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#f8f9fa")),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#e2e8f0")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#edf2f7")),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
    ]))
    elements.append(info_table)
    elements.append(Spacer(1, 15))
    
    pay_data = [
        [Paragraph("Description", header_left_style), Paragraph("Earnings (Rp)", header_right_style), Paragraph("Deductions (Rp)", header_right_style)],
        [Paragraph("Basic Salary (Gaji Pokok)", normal_style), Paragraph(f"Rp {base:,.0f}", right_normal_style), Paragraph("—", right_normal_style)],
        [Paragraph("Allowances (Tunjangan)", normal_style), Paragraph(f"Rp {allowance:,.0f}", right_normal_style), Paragraph("—", right_normal_style)],
        [Paragraph("Deductions (Potongan/BPJS/PPH)", normal_style), Paragraph("—", right_normal_style), Paragraph(f"Rp {deduction:,.0f}", right_normal_style)],
        [Paragraph("TOTAL", bold_style), Paragraph(f"Rp {(base + allowance):,.0f}", right_bold_style), Paragraph(f"Rp {deduction:,.0f}", right_bold_style)],
        [Paragraph("TAKE HOME PAY / NET SALARY", bold_style), Paragraph(f"<font color='#2b6cb0' size='12'><b>Rp {net:,.0f}</b></font>", net_pay_style), ""],
    ]
    
    pay_table = Table(pay_data, colWidths=[260, 130, 130])
    pay_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1e3a5f")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e0")),
        ('SPAN', (1,-1), (2,-1)),
        ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor("#ebf8ff")),
        ('BACKGROUND', (0,-2), (-1,-2), colors.HexColor("#f7fafc")),
        ('TOPPADDING', (0,0), (-1,-1), 7),
        ('BOTTOMPADDING', (0,0), (-1,-1), 7),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    elements.append(pay_table)
    elements.append(Spacer(1, 25))
    
    note_p = Paragraph("<font size='8' color='#718096'>Dokumen ini sah dan diterbitkan secara otomatis oleh HRIS PT Mitsindo tanpa tanda tangan basah.</font>", subtitle_style)
    elements.append(note_p)
    
    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()


# ===========================================================================
# LEAVE EXPORTS
# ===========================================================================

LEAVE_SELECT = """
    SELECT
        COALESCE(e.full_name, '') AS employee,
        COALESCE(lt.name, '') AS leave_type,
        lv.start_date,
        lv.end_date,
        CAST((JULIANDAY(lv.end_date) - JULIANDAY(lv.start_date)) + 1 AS INTEGER) AS days,
        lv.reason,
        lv.status
    FROM leaves lv
    LEFT JOIN employees e ON lv.employee_id = e.id
    LEFT JOIN leave_types lt ON lv.leave_type_id = lt.id
"""

LEAVE_HEADERS = [
    "Employee", "Leave Type", "Start Date", "End Date",
    "Days", "Reason", "Status",
]


def export_leave_excel(db=None, filters=None):
    """Export leaves to Excel. Returns bytes."""
    conn, opened = _get_db(db)
    where, params = _where_clause(filters)
    query = f"{LEAVE_SELECT} {where} ORDER BY lv.start_date DESC"
    raw = conn.execute(query, params).fetchall()
    _close_db(conn, opened)
    rows = [
        [r["employee"], r["leave_type"], r["start_date"], r["end_date"],
         r["days"], r["reason"], r["status"]]
        for r in raw
    ]
    return _build_workbook(LEAVE_HEADERS, rows)


# ===========================================================================
# OVERTIME EXPORTS
# ===========================================================================

OVERTIME_SELECT = """
    SELECT
        COALESCE(e.full_name, '') AS employee,
        o.date,
        o.hours,
        o.reason,
        o.status
    FROM overtime o
    LEFT JOIN employees e ON o.employee_id = e.id
"""

OVERTIME_HEADERS = ["Employee", "Date", "Hours", "Reason", "Status"]


def export_overtime_excel(db=None, filters=None):
    """Export overtime to Excel. Returns bytes."""
    conn, opened = _get_db(db)
    where, params = _where_clause(filters)
    query = f"{OVERTIME_SELECT} {where} ORDER BY o.date DESC"
    raw = conn.execute(query, params).fetchall()
    _close_db(conn, opened)
    rows = [
        [r["employee"], r["date"], r["hours"], r["reason"], r["status"]]
        for r in raw
    ]
    return _build_workbook(OVERTIME_HEADERS, rows)


# ===========================================================================
# ASSETS EXPORTS
# ===========================================================================

ASSETS_SELECT = """
    SELECT
        a.name,
        a.category,
        a.serial_number,
        a.purchase_date,
        a.purchase_price AS cost,
        a.condition,
        COALESCE(e.full_name, '') AS assigned_to,
        a.status
    FROM assets a
    LEFT JOIN employees e ON a.assigned_to = e.id
"""

ASSETS_HEADERS = [
    "Name", "Category", "Serial Number", "Purchase Date",
    "Cost", "Condition", "Assigned To", "Status",
]


def export_assets_excel(db=None, filters=None):
    """Export assets to Excel. Returns bytes."""
    conn, opened = _get_db(db)
    where, params = _where_clause(filters)
    query = f"{ASSETS_SELECT} {where} ORDER BY a.id"
    raw = conn.execute(query, params).fetchall()
    _close_db(conn, opened)
    rows = [
        [r["name"], r["category"], r["serial_number"], r["purchase_date"],
         r["cost"], r["condition"], r["assigned_to"], r["status"]]
        for r in raw
    ]
    return _build_workbook(ASSETS_HEADERS, rows, currency_cols={5})


# ===========================================================================
# PROCUREMENT EXPORTS
# ===========================================================================

PROCUREMENT_SELECT = """
    SELECT
        pr.title,
        pr.description,
        pr.estimated_cost AS cost,
        COALESCE(e.full_name, '') AS requester,
        pr.status,
        pr.created_at
    FROM procurement_requests pr
    LEFT JOIN employees e ON pr.requester_id = e.id
"""

PROCUREMENT_HEADERS = [
    "Title", "Description", "Cost", "Requester", "Status", "Created",
]


def export_procurement_excel(db=None, filters=None):
    """Export procurement requests to Excel. Returns bytes."""
    conn, opened = _get_db(db)
    where, params = _where_clause(filters)
    query = f"{PROCUREMENT_SELECT} {where} ORDER BY pr.created_at DESC"
    raw = conn.execute(query, params).fetchall()
    _close_db(conn, opened)
    rows = [
        [r["title"], r["description"], r["cost"], r["requester"],
         r["status"], r["created_at"]]
        for r in raw
    ]
    return _build_workbook(PROCUREMENT_HEADERS, rows, currency_cols={3})


# ===========================================================================
# PERFORMANCE EXPORTS
# ===========================================================================

PERFORMANCE_SELECT = """
    SELECT
        COALESCE(e.full_name, '') AS employee,
        pr.period,
        pr.review_type AS type,
        pr.score,
        pr.comments
    FROM performance_reviews pr
    LEFT JOIN employees e ON pr.employee_id = e.id
"""

PERFORMANCE_HEADERS = ["Employee", "Period", "Type", "Score", "Comments"]


def export_performance_excel(db=None, filters=None):
    """Export performance reviews to Excel. Returns bytes."""
    conn, opened = _get_db(db)
    where, params = _where_clause(filters)
    query = f"{PERFORMANCE_SELECT} {where} ORDER BY pr.period DESC"
    raw = conn.execute(query, params).fetchall()
    _close_db(conn, opened)
    rows = [
        [r["employee"], r["period"], r["type"], r["score"], r["comments"]]
        for r in raw
    ]
    return _build_workbook(PERFORMANCE_HEADERS, rows)


# ===========================================================================
# REIMBURSEMENT EXPORTS
# ===========================================================================

REIMBURSEMENT_SELECT = """
    SELECT
        COALESCE(e.full_name, '') AS employee,
        rb.category,
        rb.amount,
        rb.description,
        rb.status
    FROM reimbursement rb
    LEFT JOIN employees e ON rb.employee_id = e.id
"""

REIMBURSEMENT_HEADERS = [
    "Employee", "Category", "Amount", "Description", "Status",
]


def export_reimbursement_excel(db=None, filters=None):
    """Export reimbursement claims to Excel. Returns bytes."""
    conn, opened = _get_db(db)
    where, params = _where_clause(filters)
    query = f"{REIMBURSEMENT_SELECT} {where}"
    raw = conn.execute(query, params).fetchall()
    _close_db(conn, opened)
    rows = [
        [r["employee"], r["category"], r["amount"], r["description"], r["status"]]
        for r in raw
    ]
    return _build_workbook(REIMBURSEMENT_HEADERS, rows, currency_cols={3})


# ===========================================================================
# TRAINING EXPORTS
# ===========================================================================

TRAINING_SELECT = """
    SELECT
        t.title,
        t.trainer,
        t.start_date,
        t.end_date,
        t.capacity,
        t.status,
        (SELECT COUNT(*) FROM training_enrollments te WHERE te.training_id = t.id) AS enrolled
    FROM training t
"""

TRAINING_HEADERS = [
    "Title", "Trainer", "Start Date", "End Date",
    "Capacity", "Status", "Enrolled",
]


def export_training_excel(db=None, filters=None):
    """Export training programs to Excel. Returns bytes."""
    conn, opened = _get_db(db)
    where, params = _where_clause(filters)
    query = f"{TRAINING_SELECT} {where} ORDER BY t.start_date DESC"
    raw = conn.execute(query, params).fetchall()
    _close_db(conn, opened)
    rows = [
        [r["title"], r["trainer"], r["start_date"], r["end_date"],
         r["capacity"], r["status"], r["enrolled"]]
        for r in raw
    ]
    return _build_workbook(TRAINING_HEADERS, rows)
