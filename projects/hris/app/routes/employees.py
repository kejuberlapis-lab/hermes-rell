"""
Employee CRUD routes with role-based filtering.
"""
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from pydantic import BaseModel
from typing import Optional
from app.database import get_db
from app.services.auth_service import get_current_user
import aiosqlite

router = APIRouter(prefix="/employees", tags=["employees"])

class EmployeeCreate(BaseModel):
    employee_id_str: Optional[str] = None
    full_name: str
    email: str
    phone: Optional[str] = None
    department_id: Optional[int] = None
    position_id: Optional[int] = None
    manager_id: Optional[int] = None
    hire_date: str
    status: Optional[str] = "active"

class EmployeeUpdate(BaseModel):
    full_name: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    department_id: Optional[int] = None
    position_id: Optional[int] = None
    manager_id: Optional[int] = None
    hire_date: Optional[str] = None
    status: Optional[str] = None
    address: Optional[str] = None
    date_of_birth: Optional[str] = None
    gender: Optional[str] = None
    marital_status: Optional[str] = None
    religion: Optional[str] = None
    bank_account: Optional[str] = None
    bank_name: Optional[str] = None
    npwp: Optional[str] = None
    bpjs_ketenagakerjaan: Optional[str] = None
    bpjs_kesehatan: Optional[str] = None

def _uid(user):
    return user["user_id"] or user["id"]

def _d(row):
    return dict(row) if row else None

@router.get("")
async def list_employees(
    department_id: Optional[int] = Query(None),
    status: Optional[str] = Query(None),
    search: Optional[str] = Query(None),
    skip: int = Query(0, ge=0),
    limit: int = Query(50, ge=1, le=200),
    user=Depends(get_current_user),
    db: aiosqlite.Connection = Depends(get_db),
):
    conditions, params = [], []
    role = user.get("role", "employee")
    if role == "employee":
        uid = _uid(user)
        cursor = await db.execute("SELECT employee_id FROM users WHERE id = ?", (uid,))
        row = await cursor.fetchone()
        eid = row["employee_id"] if row and row["employee_id"] else None
        if eid:
            conditions.append("e.id = ?")
            params.append(eid)
        else:
            return {"employees": [], "total": 0}
    elif role == "manager":
        uid = _uid(user)
        cursor = await db.execute("SELECT employee_id FROM users WHERE id = ?", (uid,))
        row = await cursor.fetchone()
        if row and row["employee_id"]:
            conditions.append("(e.manager_id = ? OR e.id = ?)")
            params.extend([row["employee_id"], row["employee_id"]])

    if department_id:
        conditions.append("e.department_id = ?")
        params.append(department_id)
    if status:
        conditions.append("e.status = ?")
        params.append(status)
    if search:
        conditions.append("(e.full_name LIKE ? OR e.email LIKE ?)")
        s = f"%{search}%"
        params.extend([s, s])

    where = " WHERE " + " AND ".join(conditions) if conditions else ""
    q = f"""SELECT e.*, p.title as position_title, d.name as department_name
            FROM employees e
            LEFT JOIN positions p ON e.position_id = p.id
            LEFT JOIN departments d ON e.department_id = d.id
            {where} ORDER BY e.id LIMIT ? OFFSET ?"""
    params.extend([limit, skip])
    cursor = await db.execute(q, params)
    rows = await cursor.fetchall()

    count_q = f"SELECT COUNT(*) FROM employees e{where}"
    count_cursor = await db.execute(count_q, params[:-2])
    total = (await count_cursor.fetchone())[0]
    return {"employees": [_d(r) for r in rows], "total": total}

@router.get("/{emp_id}")
async def get_employee(emp_id: int, user=Depends(get_current_user), db: aiosqlite.Connection = Depends(get_db)):
    cursor = await db.execute(
        """SELECT e.*, p.title as position_title, d.name as department_name
           FROM employees e
           LEFT JOIN positions p ON e.position_id = p.id
           LEFT JOIN departments d ON e.department_id = d.id
           WHERE e.id = ?""", (emp_id,)
    )
    emp = await cursor.fetchone()
    if not emp:
        raise HTTPException(status_code=404, detail="Employee not found")
    if user["role"] == "employee":
        uid = _uid(user)
        c2 = await db.execute("SELECT employee_id FROM users WHERE id = ?", (uid,))
        own = await c2.fetchone()
        if not own or own["employee_id"] != emp_id:
            raise HTTPException(status_code=403, detail="Access denied")
    return _d(emp)

@router.post("", status_code=201)
async def create_employee(emp: EmployeeCreate, user=Depends(get_current_user), db: aiosqlite.Connection = Depends(get_db)):
    if user["role"] not in ("super_admin", "manager"):
        raise HTTPException(status_code=403, detail="Insufficient permissions")
    from datetime import datetime as _dt
    # Auto-generate employee_id_str if not provided
    eid_str = emp.employee_id_str or f"MVP-{_dt.now().strftime('%Y%m%d')}-{str(_dt.now().timestamp()).split('.')[1][-4:]}"
    try:
        cursor = await db.execute(
            """INSERT INTO employees (employee_id_str, full_name, email, phone,
               department_id, position_id, manager_id, hire_date, status)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (eid_str, emp.full_name, emp.email, emp.phone,
             emp.department_id, emp.position_id, emp.manager_id, emp.hire_date, emp.status)
        )
        await db.commit()
        return {"id": cursor.lastrowid, "message": "Employee created"}
    except Exception as e:
        if "UNIQUE" in str(e):
            raise HTTPException(status_code=400, detail="Employee ID or email already exists")
        raise

@router.put("/{emp_id}")
async def update_employee(emp_id: int, emp: EmployeeUpdate, user=Depends(get_current_user), db: aiosqlite.Connection = Depends(get_db)):
    caller_emp_id = user.get("employee_id") or user.get("id")
    # Allow user to update their own profile, or manager/admin/director
    if caller_emp_id != emp_id and user["role"] not in ("super_admin", "director", "manager"):
        raise HTTPException(status_code=403, detail="Insufficient permissions")
    
    fields = {k: v for k, v in emp.dict().items() if v is not None}
    
    # If standard employee editing their own profile, protect core HR fields like role, dept, pos, status
    if user["role"] == "employee" and caller_emp_id == emp_id:
        protected_fields = ["department_id", "position_id", "manager_id", "hire_date", "status"]
        for pf in protected_fields:
            fields.pop(pf, None)
            
    if not fields:
        raise HTTPException(status_code=400, detail="No fields to update")
    set_clause = ", ".join(f"{k} = ?" for k in fields)
    await db.execute(f"UPDATE employees SET {set_clause} WHERE id = ?", list(fields.values()) + [emp_id])
    await db.commit()
    return {"message": "Employee updated successfully"}

@router.delete("/{emp_id}")
async def delete_employee(emp_id: int, user=Depends(get_current_user), db: aiosqlite.Connection = Depends(get_db)):
    if user["role"] not in ("super_admin", "director"):
        raise HTTPException(status_code=403, detail="Admin or Director access required")
    try:
        await db.execute("DELETE FROM employees WHERE id = ?", (emp_id,))
        await db.commit()
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Cannot delete: employee has related records. ({str(e)[:100]})")
    return {"message": "Employee deleted"}

@router.post("/import")
async def import_employees(
    file: UploadFile = File(...),
    user=Depends(get_current_user),
    db: aiosqlite.Connection = Depends(get_db)
):
    if user["role"] not in ("super_admin", "director", "hr_admin"):
        raise HTTPException(status_code=403, detail="HR Admin, Director or above required")

    content = await file.read()

    if file.filename.endswith('.csv'):
        import csv, io
        text = content.decode('utf-8')
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)
    elif file.filename.endswith(('.xlsx', '.xls')):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            rows.append(dict(zip(headers, row)))
    else:
        raise HTTPException(status_code=400, detail="Only CSV and XLSX files supported")

    imported = 0
    errors = []
    for i, row in enumerate(rows):
        try:
            full_name = row['full_name'] or row['Full Name'] or row.get('name', '')
            email = row['email'] or row.get('Email', '')
            if not full_name or not email:
                errors.append(f"Row {i+2}: missing name or email")
                continue

            phone = row['phone'] or row.get('Phone', '')
            dept = row['department'] or row.get('Department', 'General')
            position = row['position'] or row.get('Position', 'Staff')
            status = row['status'] or row.get('Status', 'active')

            # Get or create department
            await db.execute("INSERT OR IGNORE INTO departments (name) VALUES (?)", (dept,))
            cur = await db.execute("SELECT id FROM departments WHERE name = ?", (dept,))
            dept_row = await cur.fetchone()
            dept_id = dept_row[0] if dept_row else 1

            # Get or create position
            await db.execute("INSERT OR IGNORE INTO positions (name, department_id) VALUES (?, ?)", (position, dept_id))
            cur = await db.execute("SELECT id FROM positions WHERE name = ? AND department_id = ?", (position, dept_id))
            pos_row = await cur.fetchone()
            pos_id = pos_row[0] if pos_row else 1

            # Check duplicate email
            cur = await db.execute("SELECT id FROM employees WHERE email = ?", (email,))
            if await cur.fetchone():
                errors.append(f"Row {i+2}: email {email} already exists")
                continue

            await db.execute(
                "INSERT INTO employees (full_name, email, phone, department_id, position_id, status, join_date) VALUES (?, ?, ?, ?, ?, ?, date('now'))",
                (full_name, email, phone, dept_id, pos_id, status)
            )
            imported += 1
        except Exception as e:
            errors.append(f"Row {i+2}: {str(e)[:100]}")

    await db.commit()
    return {"imported": imported, "errors": errors, "total_rows": len(rows)}
